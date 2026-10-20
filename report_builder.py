import os
import docx
from docx.shared import Pt, Inches, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_base_document():
    doc = docx.Document()
    
    # 1. Поля по ГОСТ: верхнее 20 мм, нижнее 20 мм, левое 30 мм, правое 15 мм
    for section in doc.sections:
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.different_first_page_header_footer = True
        
        # Нумерация страниц в нижнем колонтитуле (начиная со страницы 2)
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = f_p.add_run()
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(11)
        f_run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        f_p._p.append(fldSimple)

    # Базовый стиль
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(14)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.first_line_indent = Mm(12.5)
    style_normal.paragraph_format.space_after = Pt(0)
    style_normal.paragraph_format.space_before = Pt(0)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    return doc

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, size=14, indent=12.5, space_after=0, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Mm(indent) if indent else 0
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
    return p

def add_heading_1(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    return p

def add_heading_2(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(12.5)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    return p

def add_heading_3(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(12.5)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    return p

def add_bullet(doc, text="", level=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Mm(12.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    bullet_symbol = "— " if level == 0 else "• "
    r = p.add_run(bullet_symbol + text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    return p

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tblPr.append(tblBorders)

def add_table(doc, title, headers, rows_data, col_widths=None):
    if title:
        tp = doc.add_paragraph()
        tp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tp.paragraph_format.first_line_indent = 0
        tp.paragraph_format.space_before = Pt(12)
        tp.paragraph_format.space_after = Pt(4)
        tp.paragraph_format.keep_with_next = True
        tr = tp.add_run(title)
        tr.font.name = 'Times New Roman'
        tr.font.size = Pt(12)
        tr.bold = False

    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(str(header_text))
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.bold = True
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Data Rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, cell_value in enumerate(row_values):
            if c_idx < len(row_cells):
                row_cells[c_idx].text = ""
                p = row_cells[c_idx].paragraphs[0]
                # If first column is short ID or number, center it
                val_str = str(cell_value)
                align = WD_ALIGN_PARAGRAPH.CENTER if (len(val_str) < 15 and c_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
                p.alignment = align
                p.paragraph_format.first_line_indent = 0
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(val_str)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10.5)
                row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width

    # Add space after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)
    sp.paragraph_format.first_line_indent = 0
    return table

def add_figure(doc, image_path, caption, width_in_inches=5.8):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.first_line_indent = 0
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=Inches(width_in_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.first_line_indent = 0
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(12)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(12)
    r_cap.italic = True
    return p_cap

def add_code_block(doc, code_text):
    for line in code_text.splitlines():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        r = p.add_run(line if line else " ")
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

def add_title_page(doc, lab_num, lab_title, discipline, role_name, student_fio, student_group="4434", teacher_fio="А. Э. Зянчурин", teacher_pos="ассистент"):
    # Верхняя шапка ГУАП
    add_p(doc, "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, bold=True, indent=0, space_after=2)
    add_p(doc, "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ АВТОНОМНОЕ ОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, indent=0, space_after=2)
    add_p(doc, "«САНКТ-ПЕТЕРБУРГСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ АЭРОКОСМИЧЕСКОГО ПРИБОРОСТРОЕНИЯ»", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, bold=True, indent=0, space_after=2)
    add_p(doc, "ИНСТИТУТ ИНФОРМАЦИОННЫХ СИСТЕМ И ЗАЩИТЫ ИНФОРМАЦИИ", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, indent=0, space_after=2)
    add_p(doc, "КАФЕДРА № 43", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, bold=True, indent=0, space_after=18)

    # Оценка и подпись преподавателя
    eval_tbl = doc.add_table(rows=2, cols=3)
    eval_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(eval_tbl)
    
    # Cells
    eval_tbl.cell(0, 0).text = "ОТЧЕТ ЗАЩИЩЕН С ОЦЕНКОЙ"
    eval_tbl.cell(0, 1).text = "ПРЕПОДАВАТЕЛЬ"
    eval_tbl.cell(0, 2).text = f"{teacher_pos} {teacher_fio}"
    eval_tbl.cell(1, 0).text = ""
    eval_tbl.cell(1, 1).text = "подпись, дата"
    eval_tbl.cell(1, 2).text = "инициалы, фамилия"

    for r in eval_tbl.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = 0
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(9.5)

    add_p(doc, "", indent=0, space_after=24)

    # Название работы
    add_p(doc, f"ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № {lab_num}", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, indent=0, space_after=6)
    add_p(doc, f"«{lab_title}»", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True, indent=0, space_after=6)
    add_p(doc, f"по дисциплине: «{discipline}»", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, indent=0, space_after=4)
    add_p(doc, f"Индивидуальный раздел: «{role_name}»", align=WD_ALIGN_PARAGRAPH.CENTER, size=12.5, italic=True, indent=0, space_after=36)

    # Таблица автора
    stud_tbl = doc.add_table(rows=2, cols=3)
    stud_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(stud_tbl)

    stud_tbl.cell(0, 0).text = "РАБОТУ ВЫПОЛНИЛ"
    stud_tbl.cell(0, 1).text = f"СТУДЕНТ гр. № {student_group}"
    stud_tbl.cell(0, 2).text = student_fio
    stud_tbl.cell(1, 0).text = ""
    stud_tbl.cell(1, 1).text = "подпись, дата"
    stud_tbl.cell(1, 2).text = "инициалы, фамилия"

    for r in stud_tbl.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = 0
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(9.5)

    add_p(doc, "", indent=0, space_after=40)
    add_p(doc, "Санкт-Петербург\n2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, indent=0)

    doc.add_page_break()

def add_assignment(doc, task_title, task_items, student_fio, student_group="4434", teacher_fio="А. Э. Зянчурин"):
    add_heading_1(doc, "ЗАДАНИЕ")
    add_p(doc, f"На индивидуальное выполнение раздела лабораторной работы студентом группы {student_group} {student_fio}.", italic=True, indent=12.5, space_after=8)
    add_p(doc, f"Тема задания: {task_title}", bold=True, indent=12.5, space_after=10)
    add_p(doc, "Перечень подлежащих исследованию, проектированию и разработке вопросов:", indent=12.5, space_after=6)
    
    for item in task_items:
        add_bullet(doc, item)

    add_p(doc, "", indent=0, space_after=18)
    
    sign_p = doc.add_paragraph()
    sign_p.paragraph_format.first_line_indent = 0
    sign_p.paragraph_format.line_spacing = 1.3
    r = sign_p.add_run(
        "Задание выдал: ____________________ / " + teacher_fio + " /\n"
        "Подпись: _________________________\n"
        "Задание принял: ___________________ / " + student_fio + " /\n"
        "Дата выдачи: «____» ___________ 2026 г."
    )
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    doc.add_page_break()

def add_toc(doc, toc_entries):
    add_heading_1(doc, "СОДЕРЖАНИЕ")
    for item in toc_entries:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        if not item.startswith("    "):
            r.bold = True
    doc.add_page_break()

def add_sources(doc, sources_list):
    add_heading_1(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    for idx, s in enumerate(sources_list):
        add_p(doc, f"{idx+1}. {s}", align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=12.5, space_after=4)

def add_appendix_heading(doc, letter, title):
    doc.add_page_break()
    add_heading_1(doc, f"ПРИЛОЖЕНИЕ {letter}")
    add_p(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, indent=0, space_after=12)
