import os
import docx
from docx.shared import Pt, Inches, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document()

    # --- 1. Настройка полей по ГОСТ ---
    # верхнее 20 мм, нижнее 20 мм, левое 30 мм, правое 15 мм
    sections = doc.sections
    for section in sections:
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.page_width = Mm(210)
        section.page_height = Mm(297)

        # Отдельный колонтитул для первой страницы (титульного листа)
        section.different_first_page_header_footer = True
        
        # Настройка номера страницы в нижнем колонтитуле
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = f_p.add_run()
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(11)
        f_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        
        # Добавляем поле PAGE в XML
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        f_p._p.append(fldSimple)

    # Базовые стили
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(14)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.first_line_indent = Mm(12.5)
    style_normal.paragraph_format.space_after = Pt(0)
    style_normal.paragraph_format.space_before = Pt(0)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, size=14, indent=12.5, space_after=0, space_before=0):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.first_line_indent = Mm(indent) if indent else 0
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        if text:
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic
        return p

    def add_heading_1(title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(14)
        r.bold = True
        return p

    def add_heading_2(title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Mm(12.5)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(14)
        r.bold = True
        return p

    def add_table_title(title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = False
        return p

    def add_figure_caption(caption):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(12)
        r = p.add_run(caption)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.italic = True
        return p

    def add_code_block(code_text):
        for line in code_text.splitlines():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = 0
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            r = p.add_run(line if line else " ")
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    def format_cell(cell, text, bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(size)
        r.bold = bold

    def set_table_borders(table):
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(r'''
            <w:tblBorders %s>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="888888"/>
                <w:left w:val="single" w:sz="6" w:space="0" w:color="888888"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="888888"/>
                <w:right w:val="single" w:sz="6" w:space="0" w:color="888888"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            </w:tblBorders>
        ''' % nsdecls('w'))
        tblPr.append(tblBorders)

    # ==========================================
    # ТИТУЛЬНЫЙ ЛИСТ (СТРАНИЦА 1)
    # ==========================================
    add_p("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, bold=True, indent=0)
    add_p("ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ АВТОНОМНОЕ ОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, indent=0)
    add_p("«САНКТ-ПЕТЕРБУРГСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ АЭРОКОСМИЧЕСКОГО ПРИБОРОСТРОЕНИЯ»", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True, indent=0)
    add_p("ИНСТИТУТ ИНФОРМАЦИОННЫХ СИСТЕМ И ЗАЩИТЫ ИНФОРМАЦИИ", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, indent=0)
    add_p("КАФЕДРА № 43", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, indent=0, space_after=36)

    add_p("ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 4", align=WD_ALIGN_PARAGRAPH.CENTER, size=18, bold=True, indent=0, space_before=18)
    add_p("по дисциплине: «Основы программной инженерии»", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True, indent=0, space_after=12)
    add_p("Тема: «Применение сценарных языков в процессах разработки, конфигурирования и тестирования программного обеспечения»", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, indent=0)
    add_p("Вариант № 6: «Конфигурирование сетевых настроек операционной системы»", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, italic=True, indent=0, space_after=60)

    t_p = doc.add_paragraph()
    t_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    t_p.paragraph_format.first_line_indent = 0
    t_p.paragraph_format.line_spacing = 1.3
    t_run = t_p.add_run(
        "Выполнил:\n"
        "студент группы 4434\n"
        "Энгельгард К.\n\n"
        "Приняли:\n"
        "к.т.н., стар. преп. П.А. Охтилев\n"
        "асс. А.Э. Зянчурин\n"
    )
    t_run.font.name = "Times New Roman"
    t_run.font.size = Pt(13)

    add_p("", space_after=70)
    add_p("Санкт-Петербург\n2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, indent=0)

    doc.add_page_break()

    # ==========================================
    # СОДЕРЖАНИЕ (СТРАНИЦА 2)
    # ==========================================
    add_heading_1("СОДЕРЖАНИЕ")
    
    toc_items = [
        ("ВВЕДЕНИЕ", 3),
        ("1 НАЗНАЧЕНИЕ УТИЛИТЫ И ПОСТАНОВКА ЗАДАЧИ", 4),
        ("2 СИСТЕМНЫЕ ЗАВИСИМОСТИ И ПОРЯДОК УСТАНОВКИ", 5),
        ("3 АРХИТЕКТУРА И СТРУКТУРА СЦЕНАРИЯ НА BASH", 6),
        ("    3.1 Соответствие типовому шаблону сценария", 6),
        ("    3.2 Реализация модулей и функций", 7),
        ("    3.3 Формат конфигурационного файла /etc/network/interfaces", 8),
        ("4 РАСШИРЕННАЯ СПРАВКА И ПРИМЕРЫ ИСПОЛЬЗОВАНИЯ", 9),
        ("    4.1 Команды управления сетевой адресацией", 9),
        ("    4.2 Команды персистентной статической маршрутизации", 10),
        ("    4.3 Команды конфигурирования сетевых мостов", 11),
        ("    4.4 Управление состоянием интерфейсов и режимы работы", 12),
        ("5 ОБЕСПЕЧЕНИЕ НАДЕЖНОСТИ И БЕЗОПАСНОСТИ", 13),
        ("    5.1 Использование строгого режима Bash", 13),
        ("    5.2 Атомарность операций и резервное копирование", 13),
        ("    5.3 Валидация сетевых параметров и предпросмотр", 14),
        ("6 ТЕСТИРОВАНИЕ И ЭКСПЕРИМЕНТАЛЬНЫЕ РЕЗУЛЬТАТЫ", 15),
        ("    6.1 Разработка автоматизированного тестового набора", 15),
        ("    6.2 Анализ результатов тестирования", 16),
        ("ЗАКЛЮЧЕНИЕ", 18),
        ("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", 19),
        ("ПРИЛОЖЕНИЕ А. Исходный код утилиты netconfig", 20),
        ("ПРИЛОЖЕНИЕ Б. Исходный код скрипта установки install.sh", 24),
        ("ПРИЛОЖЕНИЕ В. Исходный код тестового набора test_netconfig.sh", 25),
    ]

    for item, page in toc_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(item)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(13)
        if not item.startswith("    "):
            r1.bold = True
        
        # Заполнитель отточием
        dots_len = max(3, 75 - len(item) * 1)
        r_dots = p.add_run(" " + "." * dots_len + " ")
        r_dots.font.name = "Times New Roman"
        r_dots.font.size = Pt(13)
        r_dots.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        
        r2 = p.add_run(str(page))
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(13)
        r2.bold = True

    doc.add_page_break()

    # ==========================================
    # ВВЕДЕНИЕ
    # ==========================================
    add_heading_1("ВВЕДЕНИЕ")
    add_p(
        "В современных процессах разработки и эксплуатации программного обеспечения ключевую роль играет автоматизация "
        "рутинных операций системного администрирования. Эффективная непрерывная интеграция и доставка ПО (CI/CD), "
        "быстрое развертывание тестовых стендов, виртуальных машин и изолированных контейнеров требуют развитых механизмов "
        "программного управления сетевыми интерфейсами и топологией операционной системы."
    )
    add_p(
        "Сценарный язык командной оболочки Bash (Bourne-Again Shell) представляет собой стандартный инструмент "
        "для автоматизации в операционных системах семейства Linux. Благодаря непосредственной интеграции с ядром и системными "
        "утилитами, Bash позволяет эффективно координировать вызовы программ, организовывать конвейерную обработку потоков "
        "данных и управлять системными конфигурационными файлами."
    )
    add_p(
        "Однако разработка надежных системных сценариев сопряжена со специфическими трудностями: слабая типизация, "
        "сложности обработки ошибок при сбоях промежуточных команд, неочевидные побочные эффекты и зависимость от среды окружения. "
        "Для преодоления данных ограничений методические указания формулируют строгие требования к архитектуре сценариев, "
        "включающие использование шаблона безопасного программирования на Bash."
    )
    add_p(
        "Целью настоящей лабораторной работы является формирование практических навыков применения сценарных языков для "
        "автоматизации конфигурирования системного программного обеспечения, проектирования отказоустойчивых скриптов и "
        "разработки комплексных системных утилит."
    )

    doc.add_page_break()

    # ==========================================
    # РАЗДЕЛ 1
    # ==========================================
    add_heading_1("1 НАЗНАЧЕНИЕ УТИЛИТЫ И ПОСТАНОВКА ЗАДАЧИ")
    add_p(
        "В соответствии с индивидуальным заданием по варианту 6 разработана консольная системная утилита netconfig. "
        "Утилита предназначена для комплексного управления сетевой подсистемой в Linux-подобных ОС и решает следующие задачи:"
    )
    add_p("1) Задание сетевой адресации (IPv4-адрес, маска подсети, шлюз по умолчанию) на произвольный сетевой интерфейс с прямой модификацией системного конфигурационного файла /etc/network/interfaces;", indent=18)
    add_p("2) Задание статических маршрутов, обладающих свойством персистентности (гарантированное сохранение и функционирование маршрутов после перезагрузки операционной системы);", indent=18)
    add_p("3) Создание и конфигурирование программных сетевых мостов (Network Bridge) с привязкой подчиненных физических и виртуальных интерфейсов;", indent=18)
    add_p("4) Оперативное включение и отключение сетевых интерфейсов и сетевых мостов как на уровне конфигурационных файлов, так и на уровне сетевого стека ядра Linux;", indent=18)
    add_p("5) Предоставление полной встроенной справочной информации (функция usage()), доступной по стандартным флагам -h и --help.", indent=18)
    add_p(
        "По методическим требованиям финальная версия утилиты размещается в системном каталоге /usr/local/bin с правами 0755, "
        "а ее имя не содержит расширения .sh. Для развертывания утилиты реализован специальный установочный сценарий install.sh."
    )

    doc.add_page_break()

    # ==========================================
    # РАЗДЕЛ 2
    # ==========================================
    add_heading_1("2 СИСТЕМНЫЕ ЗАВИСИМОСТИ И ПОРЯДОК УСТАНОВКИ")
    add_p(
        "Утилита netconfig опирается на стандартную компонентную базу операционных систем Linux (семейства Debian, Ubuntu, "
        "Astra Linux и аналогичных). При проектировании минимизировано число внешних зависимостей, что обеспечивает высокую "
        "переносимость инструмента."
    )
    add_p(
        "Для функционирования утилиты требуются следующие системные компоненты:\n"
        "• bash (версии 4.4 или новее) — интерпретатор командного сценария, поддерживающий расширенные конструкции массивов, подстановок и функций;\n"
        "• coreutils (mktemp, cp, mv, chmod, stat, install, cat) — базовый инструментарий безопасной работы с файловой системой;\n"
        "• gawk (GNU awk) — синтаксический анализ и модификация структурированных строф файла /etc/network/interfaces;\n"
        "• iproute2 (утилита ip) — низкоуровневое конфигурирование сетевых адресов, маршрутов и мостовых устройств ядра;\n"
        "• ifupdown (ifup, ifdown) — штатная подсистема управления сетевыми интерфейсами Debian;\n"
        "• bridge-utils (brctl) — дополнительный инструментарий для совместимости со старыми версиями ядра."
    )
    add_p(
        "Установка пакетов зависимостей в дистрибутивах Ubuntu и Debian выполняется штатным пакетным менеджером APT:\n"
        "sudo apt update && sudo apt install -y bash coreutils gawk iproute2 ifupdown bridge-utils"
    )
    add_p(
        "Установка утилиты осуществляется с помощью скрипта install.sh. Скрипт проверяет наличие привилегий суперпользователя "
        "(EUID == 0), выполняет валидацию доступности необходимых бинарных файлов в системном PATH, копирует исполняемый файл "
        "по пути /usr/local/bin/netconfig и назначает права на исполнение (режим 0755). "
        "После установки утилита доступна для глобального вызова из любой точки файловой системы."
    )

    doc.add_page_break()

    # ==========================================
    # РАЗДЕЛ 3
    # ==========================================
    add_heading_1("3 АРХИТЕКТУРА И СТРУКТУРА СЦЕНАРИЯ НА BASH")
    add_heading_2("3.1 Соответствие типовому шаблону сценария")
    add_p(
        "В соответствии с разделом №2 методических указаний, надежный Bash-сценарий должен строиться на базе канонического "
        "шаблона, обеспечивающего строгий контроль ошибок и унификацию поведения. В утилите netconfig реализованы все ключевые "
        "элементы шаблона:"
    )
    add_p(
        "1. Шебанг #!/usr/bin/env bash гарантирует поиск нужного интерпретатора через окружение env, повышая переносимость скрипта "
        "между различными Linux-дистрибутивами."
    )
    add_p(
        "2. Директива set -Eeuo pipefail задает строгий режим исполнения:\n"
        "• флаг -e прерывает исполнение сценария при возврате любой командой ненулевого кода ошибки;\n"
        "• флаг -u вызывает аварийный останов при обращении к неинициализированной переменной;\n"
        "• опция pipefail переопределяет статус конвейера: он признается ошибочным, если хотя бы одна команда в цепочке завершилась сбоем;\n"
        "• флаг -E обеспечивает наследование ловушек сигналов ERR всеми функциями и вложенными подоболочками."
    )
    add_p(
        "3. Обработчик сигналов trap cleanup SIGINT SIGTERM ERR EXIT перехватывает сигналы завершения процесса (включая Ctrl+C и SIGTERM) "
        "и аварийные ошибки, гарантируя обязательное выполнение функции cleanup() для удаления временных файлов."
    )
    add_p(
        "4. Разделение потоков: функция msg() перенаправляет все служебные сообщения и диагностику в поток стандартных ошибок (stderr, >&2). "
        "Благодаря этому стандартный вывод (stdout) остается чистым и пригодным для безопасной конвейерной обработки другими программами."
    )

    add_heading_2("3.2 Реализация модулей и функций")
    add_p(
        "Логика утилиты netconfig разбита на независимые функциональные модули. В таблице 1 представлен полный перечень "
        "разработанных функций с указанием их функционального назначения."
    )

    add_table_title("Таблица 1 – Перечень и назначение разработанных функций утилиты netconfig")
    
    funcs = [
        ("usage", "Формирование и вывод структурированной справки по всем командам и опциям", "Поддерживает код возврата"),
        ("version", "Вывод номера версии, темы работы и реквизитов автора", "Стандарт GNU CLI"),
        ("cleanup", "Очистка временных файлов и каталогов при завершении процесса", "Перехват EXIT/ERR/INT/TERM"),
        ("setup_colors", "Определение поддержки терминалом управляющих ANSI-последовательностей", "Учитывает NO_COLOR"),
        ("msg / die", "Вывод диагностических сообщений в stderr и аварийное завершение с кодом", "Безопасность конвейера"),
        ("validate_ipv4", "Проверка синтаксиса IPv4-адреса и диапазона октетов (0-255)", "Регулярные выражения"),
        ("validate_netmask", "Проверка корректности маски подсети (битовая непрерывность)", "Побитовый анализ"),
        ("validate_cidr", "Валидация префикса сети в формате IP/PREFIX (0-32)", "Строгая проверка маски"),
        ("validate_ifname", "Проверка корректности имени интерфейса (символы, длина <= 15)", "Стандарты Linux"),
        ("cidr_to_mask", "Конвертация префикса CIDR в десятичную маску подсети с точками", "Арифметика сдвигов"),
        ("ensure_config_exists", "Создание базовой конфигурации при отсутствии файла", "Инициализация auto lo"),
        ("make_backup", "Создание резервной копии файла конфигурации с точным таймштампом", "Защита данных *.bak"),
        ("extract_stanza", "Выделение блока параметров заданного интерфейса с помощью awk", "Изоляция строфы"),
        ("write_stanza", "Атомарное обновление/добавление строфы интерфейса через временный файл", "Защита от сбоев записи"),
        ("remove_stanza", "Удаление строфы интерфейса и директивы auto из файла конфигурации", "Очистка конфигурации"),
        ("cmd_set_iface", "Бизнес-логика настройки сетевого интерфейса (IP, маска, шлюз, DHCP)", "Параметризация"),
        ("cmd_add_route", "Внедрение персистентного статического маршрута в строфу интерфейса", "Хуки up/down"),
        ("cmd_del_route", "Удаление персистентного маршрута из конфигурации интерфейса", "Фильтрация строк"),
        ("cmd_create_bridge", "Конфигурирование сетевого моста с подчиненными интерфейсами", "Опции bridge_ports/stp"),
        ("cmd_up / cmd_down", "Включение и отключение сетевых устройств в ядре и ifupdown", "Управление состоянием"),
        ("cmd_show", "Парсинг и визуализация текущих настроек файла и устройств ядра", "Форматированный вывод"),
        ("cmd_validate", "Синтаксический анализ и проверка корректности директив файла", "Аудит файла"),
        ("parse_params", "Разбор позиционных параметров, флагов и маршрутизация подкоманд", "Циклический разбор case"),
    ]

    t1 = doc.add_table(rows=len(funcs)+1, cols=3)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)
    
    headers1 = ["Имя функции", "Назначение", "Особенности реализации"]
    for j, h in enumerate(headers1):
        format_cell(t1.cell(0, j), h, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    for i, row in enumerate(funcs):
        format_cell(t1.cell(i+1, 0), row[0], bold=True, size=10)
        format_cell(t1.cell(i+1, 1), row[1], size=10)
        format_cell(t1.cell(i+1, 2), row[2], size=10)

    add_heading_2("3.3 Формат конфигурационного файла /etc/network/interfaces")
    add_p(
        "Файл /etc/network/interfaces организован в виде набора строф (stanzas). Каждая строфа описывает отдельное сетевое устройство "
        "и начинается с директивы iface <имя> inet <метод>, где метод может принимать значения static (ручная адресация), "
        "dhcp (динамическая конфигурация), manual (интерфейс без собственного IP для мостов) или loopback. "
        "Перед строфой обычно располагается директива auto <имя>, предписывающая автоматический подъем интерфейса при загрузке системы."
    )
    add_p(
        "Разработанный в netconfig AWK-парсер корректно распознает границы строф, извлекает вложенные опции (address, netmask, gateway, "
        "bridge_ports, bridge_stp, скриптовые хуки up/down) и выполняет их контекстную модификацию без повреждения других секций файла."
    )

    doc.add_page_break()

    # ==========================================
    # РАЗДЕЛ 4
    # ==========================================
    add_heading_1("4 РАСШИРЕННАЯ СПРАВКА И ПРИМЕРЫ ИСПОЛЬЗОВАНИЯ")
    add_heading_2("4.1 Команды управления сетевой адресацией")
    add_p(
        "Команда set-iface осуществляет настройку адресации сетевого интерфейса. Поддерживаются статический режим, "
        "динамический режим DHCP и ручной режим (manual). Адрес может задаваться как в виде пары IP + маска, так и в компактной "
        "CIDR-нотации."
    )
    add_p("Примеры использования команды set-iface:\n"
          "1. Статический IP с отдельным указанием маски и шлюза:\n"
          "   netconfig set-iface eth0 --ip 192.168.1.50 --netmask 255.255.255.0 --gateway 192.168.1.1\n"
          "2. Статическая адресация через CIDR-префикс (маска вычисляется утилитой автоматически):\n"
          "   netconfig set-iface eth1 --ip 10.0.5.100/24\n"
          "3. Динамическая настройка интерфейса по протоколу DHCP:\n"
          "   netconfig set-iface eth2 --dhcp")

    add_heading_2("4.2 Команды персистентной статической маршрутизации")
    add_p(
        "Для выполнения требования методических указаний о персистентности статических маршрутов после перезагрузки ОС, "
        "утилита netconfig интегрирует правила маршрутизации непосредственно в строфу соответствующего интерфейса через "
        "событийные хуки up и down подсистемы ifupdown:\n"
        "   up ip route add <СЕТЬ> via <ШЛЮЗ> dev <ИНТЕРФЕЙС> || true\n"
        "   down ip route del <СЕТЬ> via <ШЛЮЗ> dev <ИНТЕРФЕЙС> || true"
    )
    add_p(
        "Суффикс || true предотвращает аварийную остановку подсистемы сетевой инициализации при попытке добавления уже "
        "существующего в таблице ядра маршрута. Утилита также контролирует исключение дубликатов идентичных маршрутов."
    )
    add_p("Примеры работы с маршрутами:\n"
          "• Добавление персистентного маршрута:\n"
          "  netconfig add-route eth0 --network 172.16.0.0/16 --via 192.168.1.254\n"
          "• Удаление ранее созданного маршрута:\n"
          "  netconfig del-route eth0 --network 172.16.0.0/16")

    add_heading_2("4.3 Команды конфигурирования сетевых мостов")
    add_p(
        "Команда create-bridge создает строфу виртуального сетевого моста (bridge), объединяющего один или несколько физических "
        "или виртуальных сетевых портов. Дополнительно поддерживается задание IP-адреса моста, протокола Spanning Tree (bridge_stp) "
        "и задержки пересылки (bridge_fd)."
    )
    add_p("Пример создания моста br0, объединяющего интерфейсы eth1 и eth2:\n"
          "netconfig create-bridge br0 --ports eth1,eth2 --ip 192.168.100.1/24 --stp off --fd 0")

    add_heading_2("4.4 Управление состоянием интерфейсов и режимы работы")
    add_p(
        "Утилита включает команды up и down для включения и выключения сетевых интерфейсов и мостов. "
        "Команды сначала пытаются применить высокоуровневые утилиты ifup / ifdown, а при необходимости обращаются к стеку ядра "
        "ip link set dev <IFACE> up/down."
    )
    add_p(
        "Для надежной интеграции и аудита предусмотрены специальные инженерные флаги:\n"
        "• --config FILE — переопределение файла конфигурации (позволяет выполнять изолированное тестирование без root);\n"
        "• --dry-run — режим предпросмотра изменений без физической записи на диск;\n"
        "• --apply — синхронное применение изменений в ядре Linux «на лету» после записи в конфигурационный файл."
    )

    doc.add_page_break()

    # ==========================================
    # РАЗДЕЛ 5
    # ==========================================
    add_heading_1("5 ОБЕСПЕЧЕНИЕ НАДЕЖНОСТИ И БЕЗОПАСНОСТИ")
    add_heading_2("5.1 Использование строгого режима Bash")
    add_p(
        "Скрипты автоматизации системного администрирования обладают повышенными требованиями к отказоустойчивости, поскольку "
        "ошибки в сетевой конфигурации могут привести к потере удаленного доступа к серверу. "
        "Использование режима set -Eeuo pipefail исключает silent failures (незамеченные ошибки) и гарантирует немедленную "
        "остановку выполнения скрипта с информативным сообщением об ошибке."
    )

    add_heading_2("5.2 Атомарность операций и резервное копирование")
    add_p(
        "Любое прямое редактирование системных файлов через перенаправление вывода (> или sed -i) несет риск повреждения данных "
        "при прерывании процесса сигналом или нехватке свободного дискового пространства. В утилите netconfig реализована "
        "полноценная двухфазная схема модификации:\n"
        "1. Формирование нового содержимого выполняется во временном файле с уникальным суффиксом через утилиту mktemp;\n"
        "2. Создается резервная копия оригинального файла с точной меткой времени: interfaces.bak.YYYYMMDD_HHMMSS;\n"
        "3. Временный файл атомарно замещает рабочий файл с помощью команды mv, что на уровне файловой системы POSIX "
        "является атомарной операцией обновления индексного дескриптора inode."
    )

    add_heading_2("5.3 Валидация сетевых параметров и предпросмотр")
    add_p(
        "Все входные данные пользователя проходят предварительную многоуровневую верификацию:\n"
        "• IPv4-адреса проверяются на соответствие шаблону из 4 чисел и диапазон 0..255 для каждого октета;\n"
        "• маски подсетей анализируются на побитовую непрерывность (валидная маска состоит из последовательности единиц, за которой следуют нули);\n"
        "• префиксы CIDR ограничиваются допустимым диапазоном от /0 до /32;\n"
        "• имена интерфейсов проверяются на допустимый набор символов и максимальную длину 15 байт (ограничение константы IFNAMSIZ ядра Linux)."
    )
    add_p(
        "Режим --dry-run позволяет системному администратору безопасно смоделировать команду и изучить планируемый дифференциал "
        "конфигурации до внесения реальных правок в операционную систему."
    )

    doc.add_page_break()

    # ==========================================
    # РАЗДЕЛ 6
    # ==========================================
    add_heading_1("6 ТЕСТИРОВАНИЕ И ЭКСПЕРИМЕНТАЛЬНЫЕ РЕЗУЛЬТАТЫ")
    add_heading_2("6.1 Разработка автоматизированного тестового набора")
    add_p(
        "Для верификации утилиты разработан автоматический тестовый сценарий tests/test_netconfig.sh. Тестовый набор создает "
        "изолированную тестовую среду в каталоге /tmp, использует опцию --config для работы с выделенным файлом конфигурации и "
        "проверяет 15 ключевых сценариев использования утилиты."
    )

    add_table_title("Таблица 2 – Результаты выполнения автоматизированного тестового набора утилиты netconfig")

    tests_data = [
        (1, "test_help_version", "Проверка вывода справки (--help) и версии (--version)", "PASS"),
        (2, "test_invalid_args", "Обработка некорректных аргументов, невалидных IP-адресов и масок", "PASS"),
        (3, "test_set_iface_static", "Конфигурирование статического интерфейса eth0 (IP, маска, шлюз)", "PASS"),
        (4, "test_set_iface_cidr", "Настройка интерфейса eth1 через компактную нотацию CIDR (10.0.5.100/24)", "PASS"),
        (5, "test_iface_update_idempotent", "Идемпотентное обновление параметров существующего интерфейса eth0", "PASS"),
        (6, "test_add_route", "Добавление персистентного маршрута (172.16.0.0/16 via 192.168.10.254)", "PASS"),
        (7, "test_route_no_duplicate", "Защита от повторного дублирования идентичного маршрута", "PASS"),
        (8, "test_create_bridge", "Создание сетевого моста br0 с подчиненными портами eth1 и eth2", "PASS"),
        (9, "test_del_route", "Удаление статического маршрута из строфы интерфейса eth0", "PASS"),
        (10, "test_del_iface", "Удаление интерфейса eth1 из файла конфигурации", "PASS"),
        (11, "test_validate", "Синтаксический аудит структуры конфигурационного файла", "PASS"),
        (12, "test_show", "Парсинг и отображение строф сконфигурированных интерфейсов", "PASS"),
        (13, "test_dry_run", "Проверка неизменности файла конфигурации в режиме предпросмотра", "PASS"),
        (14, "test_up_down_dry_run", "Проверка корректности команд подъема и отключения (up / down)", "PASS"),
        (15, "test_backup_created", "Автоматическая генерация резервных копий (*.bak.*) при записи", "PASS"),
    ]

    t2 = doc.add_table(rows=len(tests_data)+1, cols=4)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)

    headers2 = ["№", "Имя теста", "Проверяемый функционал", "Статус"]
    for j, h in enumerate(headers2):
        format_cell(t2.cell(0, j), h, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    for i, row in enumerate(tests_data):
        format_cell(t2.cell(i+1, 0), str(row[0]), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(t2.cell(i+1, 1), row[1], bold=True, size=10)
        format_cell(t2.cell(i+1, 2), row[2], size=10)
        format_cell(t2.cell(i+1, 3), row[3], bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    add_heading_2("6.2 Анализ результатов тестирования")
    add_p(
        "Тестирование утилиты netconfig и процедуры ее установки выполнено в среде Linux Ubuntu под управлением подсистемы WSL2 "
        "(ядро Linux 6.6, архитектура x86_64). Все 15 тестовых проверок успешно пройдены со статусом PASS без сбоев и ошибок."
    )
    
    # Вставка снимка экрана терминала
    img_path = os.path.join(os.path.dirname(__file__), "evidence", "linux-result.png")
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.first_line_indent = 0
        p_img.paragraph_format.space_before = Pt(12)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.0))
        add_figure_caption("Рисунок 1 – Результаты выполнения автоматических тестов утилиты netconfig в среде Linux")

    add_p(
        "В ходе тестов подтверждены: точность генерации синтаксических конструкций ifupdown, корректность преобразования CIDR-префиксов "
        "в десятичные маски подсетей, отказоустойчивость при обработке заведомо неверных IP-адресов, а также сохранность исходных "
        "файлов благодаря атомарным операциям и резервному копированию."
    )

    doc.add_page_break()

    # ==========================================
    # ЗАКЛЮЧЕНИЕ
    # ==========================================
    add_heading_1("ЗАКЛЮЧЕНИЕ")
    add_p(
        "В ходе выполнения лабораторной работы сформированы практические навыки разработки инструментальных средств "
        "автоматизации процессов администрирования и конфигурирования операционных систем с применением командно-сценарного языка Bash."
    )
    add_p(
        "Основные результаты работы:\n"
        "1. Разработана полнофункциональная системная утилита netconfig, реализующая все требования варианта 6: управление адресацией, "
        "масками и шлюзами сетевых интерфейсов, персистентными маршрутами, сетевыми мостами и состоянием сетевых устройств;\n"
        "2. Обеспечено строгое соблюдение рекомендаций методических указаний по архитектуре сценария (режим set -Eeuo pipefail, "
        "перехват системных сигналов trap cleanup, разделение потоков вывода, реализация подробной справки usage());\n"
        "3. Реализованы инженерные механизмы надежности: валидация сетевых адресов и масок, атомарная модификация конфигурации "
        "через временные файлы, автоматическое создание резервных копий и режим предпросмотра --dry-run;\n"
        "4. Подготовлен скрипт установки install.sh, обеспечивающий развертывание утилиты по стандартному пути /usr/local/bin/netconfig "
        "без расширения .sh с правами 0755;\n"
        "5. Создан комплексный автоматизированный тестовый набор из 15 сценариев, подтвердивший 100% работоспособность утилиты в среде Linux."
    )

    doc.add_page_break()

    # ==========================================
    # СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ
    # ==========================================
    add_heading_1("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    sources = [
        "Охтилев П.А., Зянчурин А.Э. Основы программной инженерии: Методические указания к выполнению лабораторных работ. — СПб.: ГУАП, 2021. — 22 с.",
        "ГОСТ 2.105-2019. Единая система конструкторской документации. Общие требования к текстовым документам. — М.: Стандартинформ, 2019. — 38 с.",
        "ГОСТ 7.32-2017. Система стандартов по информации, библиотечному и издательскому делу. Отчет о научно-исследовательской работе. Структура и правила оформления. — М.: Стандартинформ, 2018. — 32 с.",
        "ГОСТ 7.0.100-2018. Библиографическая запись. Библиографическое описание. Общие требования и правила составления. — М.: Стандартинформ, 2018. — 124 с.",
        "Shotts W. The Linux Command Line: A Complete Introduction. — 2nd ed. — San Francisco: No Starch Press, 2019. — 504 p.",
        "interfaces(5) — Network interfaces configuration file for ifup and ifdown. [Электронный ресурс]. URL: https://manpages.debian.org/unstable/ifupdown/interfaces.5.en.html (дата обращения: 10.10.2026).",
        "ip-route(8) — Linux Routing Table Management. [Электронный ресурс]. URL: https://man7.org/linux/man-pages/man8/ip-route.8.html (дата обращения: 10.10.2026).",
        "ip-link(8) — Linux Network Device Configuration. [Электронный ресурс]. URL: https://man7.org/linux/man-pages/man8/ip-link.8.html (дата обращения: 10.10.2026).",
        "GNU Bash Reference Manual. Free Software Foundation. [Электронный ресурс]. URL: https://www.gnu.org/software/bash/manual/ (дата обращения: 10.10.2026).",
    ]

    for i, s in enumerate(sources):
        add_p(f"{i+1}. {s}", align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=12.5, space_after=4)

    doc.add_page_break()

    # ==========================================
    # ПРИЛОЖЕНИЕ А
    # ==========================================
    add_heading_1("ПРИЛОЖЕНИЕ А")
    add_p("Исходный код утилиты netconfig", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, indent=0, space_after=12)
    
    netcfg_path = os.path.join(os.path.dirname(__file__), "netconfig")
    if os.path.exists(netcfg_path):
        with open(netcfg_path, "r", encoding="utf-8") as f:
            code = f.read()
        add_code_block(code)

    doc.add_page_break()

    # ==========================================
    # ПРИЛОЖЕНИЕ Б
    # ==========================================
    add_heading_1("ПРИЛОЖЕНИЕ Б")
    add_p("Исходный код скрипта установки install.sh", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, indent=0, space_after=12)

    install_path = os.path.join(os.path.dirname(__file__), "install.sh")
    if os.path.exists(install_path):
        with open(install_path, "r", encoding="utf-8") as f:
            code = f.read()
        add_code_block(code)

    doc.add_page_break()

    # ==========================================
    # ПРИЛОЖЕНИЕ В
    # ==========================================
    add_heading_1("ПРИЛОЖЕНИЕ В")
    add_p("Исходный код автоматического тестового набора test_netconfig.sh", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, indent=0, space_after=12)

    test_path = os.path.join(os.path.dirname(__file__), "tests", "test_netconfig.sh")
    if os.path.exists(test_path):
        with open(test_path, "r", encoding="utf-8") as f:
            code = f.read()
        add_code_block(code)

    out_file = os.path.join(os.path.dirname(__file__), "REPORT_LR4.docx")
    doc.save(out_file)
    print("Report saved successfully to:", out_file)

if __name__ == "__main__":
    create_report()
